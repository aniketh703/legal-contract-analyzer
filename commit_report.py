import subprocess
from docx import Document

try:
    out = subprocess.check_output(['git', 'log', '-1', '--pretty=format:%H%n%an%n%ad%n%s']).decode('utf-8')
    lines = out.splitlines()
    commit_hash = lines[0] if len(lines) > 0 else ''
    author = lines[1] if len(lines) > 1 else ''
    date = lines[2] if len(lines) > 2 else ''
    message = lines[3] if len(lines) > 3 else ''

    files_out = subprocess.check_output(['git', 'diff-tree', '--no-commit-id', '--name-only', '-r', commit_hash]).decode('utf-8')
    files = [f for f in files_out.splitlines() if f]

    doc = Document()
    doc.add_heading('Commit Report', level=1)
    doc.add_paragraph(f'Commit: {commit_hash}')
    doc.add_paragraph(f'Author: {author}')
    doc.add_paragraph(f'Date: {date}')
    doc.add_paragraph(f'Message: {message}')

    doc.add_heading('Files Changed', level=2)
    for f in files:
        doc.add_paragraph(f, style='List Bullet')

    out_path = 'commit_report.docx'
    doc.save(out_path)
    print(f'Wrote {out_path}')
except Exception as e:
    print('Error generating report:', e)
